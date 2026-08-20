#!/usr/bin/env python3
"""Unit tests for reference_ingest.py + end-to-end compatibility with the
downstream generator (zotero_field_insert.insert_zotero_fields + verify).

Run:  python3 -m unittest discover -s tests
"""
import os
import sys
import json
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import reference_ingest as ri
import zotero_field_insert as zfi


class FakeZotero:
    """Records calls; returns canned responses for write_item / add_items_to_collection."""

    def __init__(self, fail=False):
        self.calls = []
        self.fail = fail
        self.next_key = ["A1B2C3D4", "E5F6G7H8", "I9J0K1L2"]

    def __call__(self, name, arguments=None):
        self.calls.append((name, arguments or {}))
        if self.fail:
            raise ConnectionError("MCP unreachable")
        if name == "write_item":
            return {"key": self.next_key.pop(0) if self.next_key else "ZZZZZZZZ"}
        if name == "add_items_to_collection":
            return {}
        if name == "get_libraries":
            return {"libraries": [{"name": "My Library"}]}
        if name == "search_library":
            return {"results": []}
        if name == "get_collections":
            return {"collections": []}
        return {}


class TestParseNumberedReference(unittest.TestCase):
    def test_bracket_style(self):
        pr = ri.parse_numbered_reference("[1] Smith J. Human emotion. Cognition. 2007.")
        self.assertEqual(pr["year"], "2007")
        self.assertEqual(pr["authors"][0]["family"], "Smith")
        self.assertIn("Human", pr["title"])

    def test_paren_and_dot_styles(self):
        for prefix in ("(2) ", "3. ", "4) "):
            pr = ri.parse_numbered_reference(prefix + "Johnson K. A title. AJournal. 2010.")
            self.assertEqual(pr["year"], "2010")
            self.assertEqual(pr["authors"][0]["family"], "Johnson")

    def test_doi_and_pmid(self):
        pr = ri.parse_numbered_reference(
            "[5] Lee M. A study. J Sci. 2020. doi:10.1000/xyz123. PMID: 31337")
        self.assertEqual(pr["doi"], "10.1000/xyz123")
        self.assertEqual(pr["pmid"], "31337")

    def test_multiword_lastname(self):
        pr = ri.parse_numbered_reference("[7] de Vries P. A paper. J. 2015.")
        self.assertEqual(pr["authors"][0]["family"], "de Vries")
        self.assertEqual(pr["base_key"], "devries2015")

    def test_two_authors_with_and(self):
        pr = ri.parse_numbered_reference("[8] Smith J and Johnson K. A title. J. 2010.")
        self.assertEqual(pr["authors"][0]["family"], "Smith")
        self.assertEqual(pr["authors"][1]["family"], "Johnson")
        self.assertIn("A title", pr["title"])

    def test_two_initial_given_name(self):
        pr = ri.parse_numbered_reference("[9] Smith JD. Title here. J Med. 2000. doi:10.1000/abc.")
        self.assertEqual(pr["authors"][0]["family"], "Smith")
        self.assertEqual(pr["authors"][0]["given"], "JD")
        self.assertEqual(pr["doi"], "10.1000/abc")
        self.assertEqual(pr["journal"], "J Med")  # DOI 不应被当成期刊

    def test_et_al_not_author(self):
        pr = ri.parse_numbered_reference("[10] Smith J, et al. Another title. J. 2015.")
        self.assertEqual(len(pr["authors"]), 1)
        self.assertEqual(pr["authors"][0]["family"], "Smith")


class TestNormalizeTitle(unittest.TestCase):
    def test_idempotent_and_strip(self):
        a = ri.normalize_title("  Human  Emotion & Cognition, 2007  ")
        self.assertEqual(a, ri.normalize_title(a))
        self.assertEqual(a, "humanemotioncognition2007")


class TestCrossRefToCsl(unittest.TestCase):
    def test_required_keys(self):
        item = {"type": "journal-article", "title": ["T"], "author": [{"family": "S", "given": "J"}],
                "issued": {"date-parts": [[2007]]}, "DOI": "10.x/y"}
        csl = ri.crossref_to_csl(item)
        for k in ("type", "title", "author", "issued"):
            self.assertIn(k, csl)
        self.assertNotIn("abstract", csl)
        self.assertEqual(csl["DOI"], "10.x/y")


class TestConfidenceScore(unittest.TestCase):
    def test_ordering(self):
        parsed = {"authors": [{"family": "Smith"}], "year": "2007",
                  "title": "Human emotion and cognition"}
        perfect = {"author": [{"family": "Smith", "given": "J"}],
                   "issued": {"date-parts": [[2007]]}, "title": "Human emotion and cognition"}
        partial = {"author": [{"family": "Smith", "given": "J"}],
                   "issued": {"date-parts": [[2008]]}, "title": "Something else entirely happens"}
        s_perfect, _ = ri.confidence_score(perfect, parsed)
        s_partial, _ = ri.confidence_score(partial, parsed)
        self.assertGreater(s_perfect, s_partial)
        self.assertGreaterEqual(s_perfect, 0.9)

    def test_partial_title_word_overlap_counts(self):
        # 修正前 _norm 去空格导致标题词重叠按整串算，部分重叠得 0 分
        parsed = {"authors": [{"family": "Smith"}], "year": "2007",
                  "title": "Human emotion and cognition"}
        same_title = {"author": [{"family": "Smith", "given": "J"}],
                      "issued": {"date-parts": [[2007]]}, "title": "Human emotion and cognition"}
        similar = {"author": [{"family": "Smith", "given": "J"}],
                   "issued": {"date-parts": [[2007]]},
                   "title": "Human emotion regulation and cognition in adults"}
        s_full, _ = ri.confidence_score(same_title, parsed)
        s_sim, _ = ri.confidence_score(similar, parsed)
        self.assertGreater(s_full, s_sim)          # 完全一致 > 部分重叠
        self.assertGreater(s_sim, 0.4)             # 部分重叠仍应给分（作者+年份+标题词重叠）


class TestParseTitleNotAuthor(unittest.TestCase):
    def test_title_starting_with_single_caps_not_misparsed_as_author(self):
        # 修正前 "A study of Y" 会被误判为作者（把 "A study of" 当姓）
        pr = ri.parse_numbered_reference("[1] A study of treatment outcomes. J Med. 2020.")
        self.assertEqual(pr["authors"], [])
        self.assertIn("study", (pr["title"] or "").lower())


class TestResolvePubmedFallback(unittest.TestCase):
    def test_low_confidence_crossref_does_not_block_pubmed(self):
        # 修正前：crossref 有命中(best!=None)就不会试 pubmed；修正后低置信度也试
        calls = []

        def fake_crossref(q):
            # 返回一个低置信度命中（作者/年份不匹配）
            return [{"type": "journal-article", "title": ["Unrelated paper"],
                     "author": [{"family": "TotallyWrong", "given": "X"}],
                     "issued": {"date-parts": [[1999]]}}]

        def fake_pubmed(q):
            calls.append("pubmed")
            return []

        orig_crossref = ri.search_crossref
        orig_pubmed = ri.search_pubmed
        ri.search_crossref = fake_crossref
        ri.search_pubmed = fake_pubmed
        try:
            refs = [(0, "smith2007", "[1] Smith J. Human emotion. Cognition. 2007.")]
            ri.resolve_metadata(refs, providers=("crossref", "pubmed"))
        finally:
            ri.search_crossref = orig_crossref
            ri.search_pubmed = orig_pubmed
        self.assertEqual(calls, ["pubmed"])  # pubmed 兜底被触发


class TestSplitInstrTextDetection(unittest.TestCase):
    """Word/Zotero 插件把一个域的指令拆成多个 <w:instrText> 时，检测也必须命中。"""

    def _body_with_split_citation(self):
        from lxml import etree
        W = zfi.W
        XMLSPACE = zfi.XMLSPACE

        def mk_run(fld=None, instr=None, text=None):
            r = etree.Element(W + 'r')
            if fld is not None:
                etree.SubElement(r, W + 'fldChar').set(W + 'fldCharType', fld)
            if instr is not None:
                it = etree.SubElement(r, W + 'instrText'); it.set(XMLSPACE, 'preserve'); it.text = instr
            if text is not None:
                t = etree.SubElement(r, W + 't'); t.set(XMLSPACE, 'preserve'); t.text = text
            return r

        body = etree.Element(W + 'body')
        p0 = etree.SubElement(body, W + 'p')
        p0.append(mk_run(fld='begin'))
        p0.append(mk_run(instr=' ADDIN ZOTERO_ITEM CSL_'))
        p0.append(mk_run(instr='CITATION {"citationID":"x"} '))
        p0.append(mk_run(fld='separate'))
        p0.append(mk_run(text='(Smith, 2007)'))
        p0.append(mk_run(fld='end'))
        p1 = etree.SubElement(body, W + 'p')
        p1.append(mk_run(text='Plain (Johnson, 2010).'))
        return body

    def test_detects_split_citation(self):
        det = zfi._detect_existing(self._body_with_split_citation())
        self.assertIn(0, det["cited_paras"])
        self.assertNotIn(1, det["cited_paras"])


class TestCheckMCPFallback(unittest.TestCase):
    def test_falls_back_when_get_libraries_missing(self):
        class Fake:
            def __call__(self, name, args=None):
                if name == "get_libraries":
                    raise ConnectionError("no such tool")
                if name == "get_collections":
                    return {"collections": []}
                return {}
        res = ri.check_zotero_mcp(zotero=Fake())
        self.assertTrue(res["ok"])
        self.assertTrue(res["configured"])


class TestCheckZoteroMCP(unittest.TestCase):
    def test_not_configured(self):
        res = ri.check_zotero_mcp(zotero=FakeZotero(fail=True))
        self.assertFalse(res["ok"])
        self.assertIn("hint", res)


class TestImportToZotero(unittest.TestCase):
    def test_itemtype_top_level_and_add(self):
        fake = FakeZotero()
        csl = {"type": "article-journal", "title": "T", "author": [{"family": "S", "given": "J"}],
               "issued": {"date-parts": [[2007]]}, "DOI": "10.x/y"}
        plan = {"s2007": {"action": "import", "csl": csl},
                "j2010": {"action": "reuse", "zotero_key": "OLD12345", "csl": csl},
                "skipme": {"action": "skip", "csl": csl}}
        imported, warnings_ = ri.import_to_zotero(plan, "COLLKEY", zotero=fake, uri_prefix="http://zotero.org/users/U/items/")
        self.assertEqual(warnings_, [])
        # write_item called with itemType as top-level param
        write_calls = [a for n, a in fake.calls if n == "write_item"]
        self.assertEqual(len(write_calls), 1)
        self.assertIn("itemType", write_calls[0])
        self.assertNotIn("itemType", write_calls[0]["fields"])
        self.assertEqual(write_calls[0]["itemType"], "journalArticle")
        # add_items_to_collection gets the new key + reused key
        add_calls = [a for n, a in fake.calls if n == "add_items_to_collection"]
        self.assertEqual(len(add_calls), 1)
        keys = add_calls[0]["itemKeys"]
        self.assertIn("A1B2C3D4", keys)
        self.assertIn("OLD12345", keys)
        self.assertNotIn("skipme", imported)
        self.assertEqual(imported["s2007"]["itemKey"], "A1B2C3D4")
        self.assertEqual(imported["j2010"]["itemKey"], "OLD12345")


class TestBuildItemMapping(unittest.TestCase):
    def test_schema(self):
        imported = {"s2007": {"itemKey": "A1B2C3D4", "uri": "http://zotero.org/users/U/items/A1B2C3D4",
                              "csl": {"type": "article-journal", "title": "T", "DOI": "10.x/y",
                                      "author": [{"family": "Smith", "given": "J"}],
                                      "issued": {"date-parts": [[2007]]}}}}
        m = ri.build_item_mapping(imported, uri_prefix="http://zotero.org/users/U/items/")
        for k in ("itemKey", "uri", "csl", "doi", "first_author", "year"):
            self.assertIn(k, m["s2007"])
        self.assertEqual(m["s2007"]["first_author"], "Smith")
        self.assertEqual(m["s2007"]["year"], "2007")


class TestEndToEndDocx(unittest.TestCase):
    def test_pipeline(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, "src.docx")
            out = os.path.join(td, "out.docx")
            doc = Document()
            doc.add_paragraph("This is a claim (Smith, 2007).")          # body para 0
            doc.add_paragraph("Another claim (Johnson, 2010).")          # body para 1
            doc.add_paragraph("[1] Smith J. Human emotion. Cognition. 2007.")  # ref para 2
            doc.add_paragraph("[2] Johnson K. A title. AJournal. 2010.")      # ref para 3
            doc.save(src)

            mapping = {
                "smith2007": {"itemKey": "A1B2C3D4",
                              "uri": "http://zotero.org/users/U/items/A1B2C3D4",
                              "csl": {"type": "article-journal", "title": "Human emotion",
                                      "author": [{"family": "Smith", "given": "J"}],
                                      "issued": {"date-parts": [[2007]]}},
                              "doi": "10.x/1", "first_author": "Smith", "year": "2007"},
                "johnson2010": {"itemKey": "E5F6G7H8",
                                "uri": "http://zotero.org/users/U/items/E5F6G7H8",
                                "csl": {"type": "article-journal", "title": "A title",
                                        "author": [{"family": "Johnson", "given": "K"}],
                                        "issued": {"date-parts": [[2010]]}},
                                "doi": "10.x/2", "first_author": "Johnson", "year": "2010"},
            }
            cit_n, bib_n = zfi.insert_zotero_fields(
                src, out, (0, 2), [2, 3], mapping,
                uri_prefix="http://zotero.org/users/U/items/")
            self.assertEqual(cit_n, 2)
            self.assertEqual(bib_n, 2)
            ok, detail = zfi.verify(out, ["A1B2C3D4", "E5F6G7H8"])
            self.assertTrue(ok, detail)
            self.assertEqual(detail["cites"], 2)
            self.assertEqual(detail["bibs"], 1)
            self.assertEqual(detail["prefs"], 1)
            self.assertEqual(detail["fldChar"].split("/")[0], detail["fldChar"].split("/")[1])


class TestResolveStyleId(unittest.TestCase):
    def test_known_aliases(self):
        self.assertEqual(ri.resolve_style_id("Nature"), ("http://www.zotero.org/styles/nature", "en-US"))
        self.assertEqual(ri.resolve_style_id("IEEE"), ("http://www.zotero.org/styles/ieee", "en-US"))
        self.assertEqual(ri.resolve_style_id("GB/T 7714"), ("http://www.zotero.org/styles/gb7714-2015-numeric", "en-US"))

    def test_unknown_returns_none(self):
        style_id, _ = ri.resolve_style_id("Some Unknown Journal")
        self.assertIsNone(style_id)


class TestMixedDoc(unittest.TestCase):
    """已含 Zotero 域代码的混合文档：重复处理不应重复插入 bib/prefs、不应重包已有引文域。"""

    def _make_plain_doc(self, td):
        from docx import Document
        src = os.path.join(td, "plain.docx")
        doc = Document()
        doc.add_paragraph("This is a claim (Smith, 2007).")
        doc.add_paragraph("Another claim (Johnson, 2010).")
        doc.add_paragraph("[1] Smith J. Human emotion. Cognition. 2007.")
        doc.add_paragraph("[2] Johnson K. A title. AJournal. 2010.")
        doc.save(src)
        return src

    def _mapping(self):
        return {
            "smith2007": {"itemKey": "A1B2C3D4", "uri": "http://zotero.org/users/U/items/A1B2C3D4",
                          "csl": {"type": "article-journal", "title": "Human emotion",
                                  "author": [{"family": "Smith", "given": "J"}],
                                  "issued": {"date-parts": [[2007]]}},
                          "doi": "10.x/1", "first_author": "Smith", "year": "2007"},
            "johnson2010": {"itemKey": "E5F6G7H8", "uri": "http://zotero.org/users/U/items/E5F6G7H8",
                            "csl": {"type": "article-journal", "title": "A title",
                                    "author": [{"family": "Johnson", "given": "K"}],
                                    "issued": {"date-parts": [[2010]]}},
                            "doi": "10.x/2", "first_author": "Johnson", "year": "2010"},
        }

    def test_detect_and_idempotent(self):
        with tempfile.TemporaryDirectory() as td:
            src = self._make_plain_doc(td)
            once = os.path.join(td, "once.docx")
            twice = os.path.join(td, "twice.docx")
            m = self._mapping()
            zfi.insert_zotero_fields(src, once, (0, 2), [2, 3], m,
                                     uri_prefix="http://zotero.org/users/U/items/")
            # 第一次处理后的文档已含域代码（prefs 隐藏段插入文首使后续段落索引 +1）
            det = zfi.detect_zotero_fields(once)
            self.assertTrue(det["has_bibliography"])
            self.assertTrue(det["has_preferences"])
            self.assertEqual(len(det["cited_paras"]), 2)  # 两个正文引文段
            # 再处理一次（混合情形：已有域代码 + 待处理纯文字），不应重复插入
            zfi.insert_zotero_fields(once, twice, (0, 2), [2, 3], m,
                                     uri_prefix="http://zotero.org/users/U/items/")
            ok, detail = zfi.verify(twice, ["A1B2C3D4", "E5F6G7H8"])
            self.assertTrue(ok, detail)
            self.assertEqual(detail["cites"], 2)   # 引文域保留，不重复
            self.assertEqual(detail["bibs"], 1)    # bibliography 不重复添加
            self.assertEqual(detail["prefs"], 1)   # preferences 不重复添加
            self.assertEqual(detail["fldChar"].split("/")[0], detail["fldChar"].split("/")[1])


class TestSameParagraphMixed(unittest.TestCase):
    """同一段落内既有已有引文域、又有新的纯文字引文：保留域，只转换纯文字。"""

    def _mapping(self):
        return {
            "smith2007": {"itemKey": "A1B2C3D4", "uri": "http://zotero.org/users/U/items/A1B2C3D4",
                          "csl": {"type": "article-journal", "title": "Human emotion",
                                  "author": [{"family": "Smith", "given": "J"}],
                                  "issued": {"date-parts": [[2007]]}},
                          "doi": "10.x/1", "first_author": "Smith", "year": "2007"},
            "johnson2010": {"itemKey": "E5F6G7H8", "uri": "http://zotero.org/users/U/items/E5F6G7H8",
                            "csl": {"type": "article-journal", "title": "A title",
                                    "author": [{"family": "Johnson", "given": "K"}],
                                    "issued": {"date-parts": [[2010]]}},
                            "doi": "10.x/2", "first_author": "Johnson", "year": "2010"},
        }

    def test_preserve_field_and_convert_plain_in_same_para(self):
        from docx import Document
        from lxml import etree
        W = zfi.W
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, "mixed.docx")
            out = os.path.join(td, "out.docx")
            doc = Document()
            para = doc.add_paragraph()
            # 模拟插件已插入的 Smith 引文域
            cj = {"citationID": "cit0",
                  "properties": {"formattedCitation": "(Smith, 2007)", "plainCitation": "(Smith, 2007)",
                                 "dontUpdate": False, "noteIndex": 0},
                  "citationItems": [{"id": 1, "uris": ["http://zotero.org/users/U/items/A1B2C3D4"],
                                     "itemData": {"type": "article-journal", "title": "Human emotion",
                                                  "author": [{"family": "Smith", "given": "J"}],
                                                  "issued": {"date-parts": [[2007]]}}}],
                  "schema": zfi.SCHEMA}
            for r in zfi._make_field_runs("ZOTERO_ITEM CSL_CITATION " + json.dumps(cj, ensure_ascii=False),
                                          "(Smith, 2007)", None):
                para._p.append(r)
            # 同一段落追加纯文字引文
            r = etree.SubElement(para._p, W + 'r')
            t = etree.SubElement(r, W + 't'); t.set(zfi.XMLSPACE, 'preserve')
            t.text = " and (Johnson, 2010)."
            doc.save(src)

            cit_n, _ = zfi.insert_zotero_fields(src, out, (0, 1), None, self._mapping(),
                                                uri_prefix="http://zotero.org/users/U/items/")
            self.assertEqual(cit_n, 1)  # 只新增 johnson，smith 已存在不计
            ok, detail = zfi.verify(out, ["A1B2C3D4", "E5F6G7H8"])
            self.assertTrue(ok, detail)
            self.assertEqual(detail["cites"], 2)  # smith(已有) + johnson(新增)
            self.assertEqual(detail["bad_json"], 0)
            self.assertEqual(detail["bad_uri"], 0)


class TestInsertTableNoTable(unittest.TestCase):
    def test_no_table_returns_zero(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, "t.docx")
            out = os.path.join(td, "o.docx")
            d = Document(); d.add_paragraph("no table"); d.save(src)
            self.assertEqual(zfi.insert_table_citations(src, out, {}, 0, 1,
                                                        uri_prefix="http://zotero.org/users/U/items/"), 0)


class TestResolveByDOI(unittest.TestCase):
    def test_doi_direct_lookup_preferred(self):
        # 参考文献自带 DOI 时，应优先按 DOI 直查（不先走模糊搜索）
        doi_calls = []
        fuzzy_calls = []

        def fake_by_doi(d):
            doi_calls.append(d)
            return [{"type": "journal-article", "title": ["Human emotion"],
                     "author": [{"family": "Smith", "given": "J"}],
                     "issued": {"date-parts": [[2007]]}, "DOI": d}]

        def fake_crossref(q):
            fuzzy_calls.append(q)
            return []

        orig_by_doi = ri.search_crossref_by_doi
        orig_crossref = ri.search_crossref
        ri.search_crossref_by_doi = fake_by_doi
        ri.search_crossref = fake_crossref
        try:
            refs = [(0, "smith2007", "[1] Smith J. Human emotion. 2007. doi:10.1000/xyz")]
            res = ri.resolve_metadata(refs, providers=("crossref",))
            self.assertEqual(doi_calls, ["10.1000/xyz"])
            self.assertEqual(fuzzy_calls, [])  # DOI 命中即不再模糊搜索
            self.assertEqual(res["smith2007"]["confidence"], 1.0)
        finally:
            ri.search_crossref_by_doi = orig_by_doi
            ri.search_crossref = orig_crossref


class TestCslFieldMapping(unittest.TestCase):
    def test_chapter_uses_bookTitle(self):
        csl = {"type": "chapter", "title": "Chap", "container-title": "Book",
               "publisher": "Pub", "issued": {"date-parts": [[2020]]}}
        fields = ri._csl_to_field_map(csl)
        self.assertEqual(fields["bookTitle"], "Book")
        self.assertNotIn("publicationTitle", fields)

    def test_journal_uses_publicationTitle(self):
        csl = {"type": "article-journal", "title": "T", "container-title": "J",
               "issued": {"date-parts": [[2020]]}}
        fields = ri._csl_to_field_map(csl)
        self.assertEqual(fields["publicationTitle"], "J")

    def test_org_creator(self):
        csl = {"type": "webpage", "title": "P", "author": [{"literal": "WHO"}]}
        creators = ri._csl_to_creators(csl)
        self.assertEqual(creators, [{"creatorType": "author", "name": "WHO"}])


class TestFullPipeline(unittest.TestCase):
    """前端 ingest -> item_mapping -> 后端 insert_zotero_fields -> verify 的端到端兼容。"""

    def test_ingest_to_insert(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, "论文.docx")
            out = os.path.join(td, "论文_zotero.docx")
            d = Document()
            d.add_paragraph("Claims (Smith, 2007) and (Jones, 2010) here.")
            d.add_paragraph("[1] Smith J. Human emotion. Cognition. 2007.")
            d.add_paragraph("[2] Jones K. A title. AJournal. 2010.")
            d.save(src)

            # 前端：fake MCP + 打桩网络
            class Fake:
                def __init__(self): self.created = []
                def __call__(self, name, args=None):
                    args = args or {}
                    if name == "get_libraries": return [{"libraryID": 1}]
                    if name == "get_collections": return []
                    if name == "search_library": return {"results": []}
                    if name == "write_item":
                        k = ["A1B2C3D4", "E5F6G7H8"][len(self.created)]
                        self.created.append(k)
                        return {"key": k}
                    if name == "add_items_to_collection": return {}
                    return {}
            fake = Fake()
            orig_cr = ri.search_crossref
            orig_doi = ri.search_crossref_by_doi
            orig_pm = ri.search_pubmed
            ri.search_crossref_by_doi = lambda doi: []
            ri.search_crossref = lambda q, rows=5: [
                {"type": "journal-article", "title": [t], "author": [{"family": a, "given": "J"}],
                 "issued": {"date-parts": [[y]]}, "DOI": f"10.x/{a.lower()}"}
                for (a, t, y) in ([("Smith", "Human emotion", 2007)] if "Smith" in q else
                                  [("Jones", "A title", 2010)] if "Jones" in q else [])
            ]
            ri.search_pubmed = lambda q, rows=5: []
            try:
                mapping, stats = ri.ingest(src, "COLLKEY", out=os.path.join(td, "item_mapping.json"), zotero=fake)
            finally:
                ri.search_crossref_by_doi = orig_doi
                ri.search_crossref = orig_cr
                ri.search_pubmed = orig_pm
            self.assertEqual(len(mapping), 2)
            self.assertEqual(mapping["smith2007"]["first_author"], "Smith")
            self.assertEqual(mapping["jones2010"]["year"], "2010")

            # 后端：用前端产出的 mapping 插入
            import zotero_field_insert as zfi
            cit_n, bib_n = zfi.insert_zotero_fields(
                src, out, (0, 1), [1, 2], mapping,
                uri_prefix="http://zotero.org/users/U/items/")
            self.assertEqual(cit_n, 2)
            ok, detail = zfi.verify(out, [v["itemKey"] for v in mapping.values()])
            self.assertTrue(ok, detail)
            self.assertEqual(detail["cites"], 2)
            self.assertEqual(detail["bibs"], 1)


class TestVerifyZeroMiss(unittest.TestCase):
    """verify 传入 src + item_mapping 时的"引文零遗漏 + 显示文本逐字一致"校验。"""

    def _mapping(self):
        return {
            "smith2007": {"itemKey": "AA", "uri": "http://zotero.org/users/U/items/AA",
                          "csl": {"type": "article-journal", "title": "Human emotion",
                                  "author": [{"family": "Smith", "given": "J"}],
                                  "issued": {"date-parts": [[2007]]}},
                          "doi": "", "first_author": "Smith", "year": "2007"},
            "jones2010": {"itemKey": "BB", "uri": "http://zotero.org/users/U/items/BB",
                          "csl": {"type": "article-journal", "title": "A title",
                                  "author": [{"family": "Jones", "given": "K"}],
                                  "issued": {"date-parts": [[2010]]}},
                          "doi": "", "first_author": "Jones", "year": "2010"},
        }

    def test_catches_missed_citation(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, "s.docx"); out = os.path.join(td, "o.docx")
            d = Document()
            d.add_paragraph("Claims (Smith, 2007) here.")
            d.add_paragraph("And (Jones, 2010) here.")   # 不在 body_para_range，应被漏抓
            d.add_paragraph("[1] Smith J. Human emotion. Cognition. 2007.")
            d.add_paragraph("[2] Jones K. A title. AJournal. 2010.")
            d.save(src)
            m = self._mapping()
            zfi.insert_zotero_fields(src, out, (0, 1), [2, 3], m,
                                     uri_prefix="http://zotero.org/users/U/items/")
            # 仅结构校验：通过（1 条引文被转换）
            ok_struct, det = zfi.verify(out, ["AA", "BB"])
            self.assertTrue(ok_struct)
            self.assertEqual(det["cites"], 1)
            # 全量校验：应标记漏抓 Jones
            ok_full, det2 = zfi.verify(out, ["AA", "BB"], src_docx=src, item_mapping=m)
            self.assertFalse(ok_full)
            self.assertIn("(Jones, 2010)", det2["missed_citations"])
            self.assertNotIn("(Smith, 2007)", det2["missed_citations"])

    def test_no_miss_when_all_converted(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, "s.docx"); out = os.path.join(td, "o.docx")
            d = Document()
            d.add_paragraph("Claims (Smith, 2007) and (Jones, 2010) here.")
            d.add_paragraph("[1] Smith J. Human emotion. Cognition. 2007.")
            d.add_paragraph("[2] Jones K. A title. AJournal. 2010.")
            d.save(src)
            m = self._mapping()
            zfi.insert_zotero_fields(src, out, (0, 1), [1, 2], m,
                                     uri_prefix="http://zotero.org/users/U/items/")
            ok_full, det2 = zfi.verify(out, ["AA", "BB"], src_docx=src, item_mapping=m)
            self.assertTrue(ok_full, det2)
            self.assertEqual(det2["missed_citations"], [])


class TestCleanCrossrefText(unittest.TestCase):
    """CrossRef 标题里的 HTML 标记/实体必须剥除（否则原样进 Zotero 标题）。"""

    def test_html_stripped_from_title_and_journal(self):
        item = {"type": "journal-article",
                "title": ["Late <i>N</i>&#8209;acetylcysteine treatment"],
                "container-title": ["<sub>Hippocampus</sub>"],
                "author": [{"family": "Lanté", "given": "Fabien"}],
                "issued": {"date-parts": [[2008]]}, "DOI": "10.x/y"}
        csl = ri.crossref_to_csl(item)
        self.assertEqual(csl["title"], "Late N‑acetylcysteine treatment")
        self.assertEqual(csl["container-title"], "Hippocampus")

    def test_plain_text_unchanged(self):
        csl = ri.crossref_to_csl({"type": "journal-article", "title": ["Plain title"]})
        self.assertEqual(csl["title"], "Plain title")


class TestDedupDoiConfirm(unittest.TestCase):
    """minimal/standard 搜索结果不含 DOI 字段：DOI 匹配必须走 get_item_details 确认。

    旧实现拿 minimal 结果直接比 it["DOI"] 恒为空 -> 漏检 -> 重复导入。"""

    def _fake(self, doi_in_lib=True, title_differs=True):
        class Fake:
            def __init__(self): self.detail_calls = []
            def __call__(self, name, args=None):
                args = args or {}
                if name == "search_library":
                    q = args.get("q", "")
                    if q.startswith('"10.'):  # DOI 全文搜索
                        return {"results": [{"key": "LIBITEM1", "title": "Some title"}]}
                    return {"results": []}  # 标题搜索无命中（大小写/连字符差异场景）
                if name == "get_item_details":
                    self.detail_calls.append(args.get("itemKey"))
                    return {"key": args.get("itemKey"),
                            "title": "Some title",
                            "DOI": "10.1000/xyz" if doi_in_lib else "10.other/zzz"}
                return {}
        return Fake()

    def test_doi_confirmed_via_details(self):
        fake = self._fake()
        resolved = {"a2007": {"csl": {"title": "Totally different-looking title",
                                      "DOI": "10.1000/xyz"}}}
        rep = ri.dedup_check(resolved, zotero=fake)
        self.assertEqual(rep["a2007"]["status"], "exact")
        self.assertEqual(rep["a2007"]["existing"]["key"], "LIBITEM1")
        self.assertEqual(fake.detail_calls, ["LIBITEM1"])

    def test_doi_mismatch_not_dup(self):
        fake = self._fake(doi_in_lib=False)
        resolved = {"a2007": {"csl": {"title": "Totally different-looking title",
                                      "DOI": "10.1000/xyz"}}}
        rep = ri.dedup_check(resolved, zotero=fake)
        self.assertEqual(rep["a2007"]["status"], "none")

    def test_audit_imported_by_doi_counts(self):
        fake = self._fake()
        resolved = {"a2007": {"csl": {"title": "t", "DOI": "10.1000/xyz"}}}
        rep = ri.audit_imported_by_doi(resolved, zotero=fake)
        self.assertEqual(rep["a2007"]["hits"], 1)
        self.assertEqual(rep["a2007"]["keys"], ["LIBITEM1"])


class TestImportEmptyCreatorsAndPrecheck(unittest.TestCase):
    """空 creators 必须整个省略（Zotero 拒绝空列表）；precheck 幂等导入。"""

    def test_empty_creators_omitted(self):
        fake = FakeZotero()
        csl = {"type": "report", "title": "EU Clinical Trial 2023-0000-00",
               "issued": {"date-parts": [[2023]]}}
        ri.import_to_zotero({"euct2023": {"action": "import", "csl": csl}}, None,
                            zotero=fake, uri_prefix="http://zotero.org/users/U/items/", precheck=False)
        write_calls = [a for n, a in fake.calls if n == "write_item"]
        self.assertEqual(len(write_calls), 1)
        self.assertNotIn("creators", write_calls[0])

    def test_precheck_reuses_existing(self):
        class Fake(FakeZotero):
            def __call__(self, name, arguments=None):
                if name == "search_library":
                    return {"results": [{"key": "EXISTING1", "title": "Same title exactly"}]}
                return super().__call__(name, arguments)
        fake = Fake()
        csl = {"type": "article-journal", "title": "Same title exactly",
               "author": [{"family": "S", "given": "J"}],
               "issued": {"date-parts": [[2007]]}, "DOI": "10.x/y"}
        imported, warnings_ = ri.import_to_zotero(
            {"s2007": {"action": "import", "csl": csl}}, None, zotero=fake,
            uri_prefix="http://zotero.org/users/U/items/", precheck=True)
        self.assertEqual(warnings_, [])
        self.assertEqual(imported["s2007"]["itemKey"], "EXISTING1")
        self.assertTrue(imported["s2007"].get("precheck_reused"))
        write_calls = [a for n, a in fake.calls if n == "write_item"]
        self.assertEqual(write_calls, [])  # 没有新建


class TestApplyLibraryData(unittest.TestCase):
    """以库为准：复用条目时 mapping 的 csl 应用库内元数据重建。"""

    def test_rebuild_from_details(self):
        class Fake:
            def __call__(self, name, args=None):
                if name == "get_item_details":
                    return {"key": args.get("itemKey"),
                            "itemType": "journalArticle",
                            "title": "Library-style Title: subtitle",
                            "creators": [{"firstName": "Gloria B.", "lastName": "Choi",
                                          "creatorType": "author"}],
                            "date": "2016-02-26",
                            "publicationTitle": "Science",
                            "volume": "351", "issue": "6276", "pages": "933-939",
                            "DOI": "10.1126/science.aad0314", "url": ""}
                return {}
        mapping = {"choi2016": {"itemKey": "UCI3CXBV", "uri": "u/UCI3CXBV",
                                "csl": {"type": "article-journal", "title": "CrossRef TITLE",
                                        "author": [{"family": "Choi", "given": "G B"}],
                                        "issued": {"date-parts": [[2016]]}},
                                "doi": "10.1126/science.aad0314",
                                "first_author": "Choi", "year": "2016"}}
        mp, mism = ri.apply_library_data(mapping, zotero=Fake())
        self.assertEqual(mism, [])
        c = mp["choi2016"]["csl"]
        self.assertEqual(c["title"], "Library-style Title: subtitle")
        self.assertEqual(c["author"][0]["given"], "Gloria B.")
        self.assertEqual(c["issued"], {"date-parts": [[2016, 2, 26]]})
        self.assertEqual(c["container-title"], "Science")
        self.assertEqual(c["page"], "933-939")
        # doi 锚点保持不变（用于正文匹配）
        self.assertEqual(mp["choi2016"]["doi"], "10.1126/science.aad0314")

    def test_doi_mismatch_reported(self):
        class Fake:
            def __call__(self, name, args=None):
                if name == "get_item_details":
                    return {"key": args.get("itemKey"), "title": "T",
                            "DOI": "10.wrong/zzz", "date": "2007",
                            "creators": [{"firstName": "J", "lastName": "S", "creatorType": "author"}]}
                return {}
        mapping = {"s2007": {"itemKey": "ABCD1234", "uri": "u/ABCD1234",
                             "csl": {"type": "article-journal", "title": "T",
                                     "author": [{"family": "S", "given": "J"}],
                                     "issued": {"date-parts": [[2007]]}},
                             "doi": "10.right/yy", "first_author": "S", "year": "2007"}}
        mp, mism = ri.apply_library_data(mapping, zotero=Fake())
        self.assertEqual(mism, [("s2007", "10.right/yy", "10.wrong/zzz")])


class TestFindDoiCitations(unittest.TestCase):
    """DOI 组括号引文 + 精确文本括号引文匹配器。"""

    def setUp(self):
        self.doi2key = {"10.1234/ax": "ax", "10.5678/by": "by"}
        self.tmap = {"FDA label, minocycline": ["fdamino"], "NCT01329991": ["nct1"]}

    def test_doi_prefixed_group(self):
        spans = zfi.find_doi_citations("(doi:10.1234/ax; doi:10.5678/by)", self.doi2key)
        self.assertEqual(len(spans), 1)
        self.assertEqual(spans[0][3], ["ax", "by"])

    def test_bare_doi_group(self):
        spans = zfi.find_doi_citations("claim (10.1234/ax; 10.5678/by) here", self.doi2key)
        self.assertEqual(len(spans), 1)
        self.assertEqual(spans[0][2], "(10.1234/ax; 10.5678/by)")

    def test_unknown_doi_skips_whole_group(self):
        spans = zfi.find_doi_citations("(10.1234/ax; 10.unknown/z)", self.doi2key)
        self.assertEqual(spans, [])

    def test_text_map_exact_match(self):
        spans = zfi.find_doi_citations("warning (FDA label, minocycline) given", {}, self.tmap)
        self.assertEqual(len(spans), 1)
        self.assertEqual(spans[0][3], ["fdamino"])

    def test_non_citation_paren_ignored(self):
        spans = zfi.find_doi_citations("(e.g., Smith, 2007) and (www.figdraw.com)", {}, self.tmap)
        self.assertEqual(spans, [])

    def test_standalone_doi(self):
        hit = zfi.match_standalone_doi(" 10.1234/ax; ", self.doi2key)
        self.assertEqual(hit, ("ax", " 10.1234/ax; "))
        self.assertIsNone(zfi.match_standalone_doi("10.1234/ax more text", self.doi2key))

    def test_end_to_end_doi_placeholder_doc(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, "s.docx"); out = os.path.join(td, "o.docx")
            d = Document()
            d.add_paragraph("Claims (doi:10.1234/ax; 10.5678/by) and (10.1234/ax) here.")
            d.add_paragraph("Reference list:")
            d.save(src)
            mapping = {
                "ax": {"itemKey": "AAAAAAAA", "uri": "http://zotero.org/users/U/items/AAAAAAAA", "doi": "10.1234/ax",
                       "csl": {"type": "article-journal", "title": "A",
                               "author": [{"family": "A", "given": "X"}],
                               "issued": {"date-parts": [[2007]]}}},
                "by": {"itemKey": "BBBBBBBB", "uri": "http://zotero.org/users/U/items/BBBBBBBB", "doi": "10.5678/by",
                       "csl": {"type": "article-journal", "title": "B",
                               "author": [{"family": "B", "given": "Y"}],
                               "issued": {"date-parts": [[2010]]}}},
            }
            cit_n, _ = zfi.insert_zotero_fields(src, out, (0, 1), [1], mapping, uri_prefix="http://zotero.org/users/U/items/")
            self.assertEqual(cit_n, 2)
            ok, det = zfi.verify(out, ["AAAAAAAA", "BBBBBBBB"])
            self.assertTrue(ok, det)
            self.assertEqual(det["cites"], 2)


class TestProofErrDoesNotSplitCitation(unittest.TestCase):
    """Word 的 <w:proofErr> 夹在括号引文组中间时，不得因元素切分而整组漏转。"""

    def test_prooferr_inside_group(self):
        from docx import Document
        from lxml import etree
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, "s.docx"); out = os.path.join(td, "o.docx")
            d = Document()
            p = d.add_paragraph()
            # 文本 "(Smith, 2007)" 被拆成 3 个 run，中间夹 proofErr
            r1 = etree.SubElement(p._p, zfi.W + 'r')
            t1 = etree.SubElement(r1, zfi.W + 't'); t1.set(zfi.XMLSPACE, 'preserve'); t1.text = "(Smith, "
            pe = etree.SubElement(p._p, zfi.W + 'proofErr'); pe.set(zfi.W + 'type', 'spellStart')
            r2 = etree.SubElement(p._p, zfi.W + 'r')
            t2 = etree.SubElement(r2, zfi.W + 't'); t2.set(zfi.XMLSPACE, 'preserve'); t2.text = "2007)"
            d.add_paragraph("[1] Smith J. Human emotion. Cognition. 2007.")
            d.save(src)
            mapping = {"smith2007": {"itemKey": "AA", "uri": "u/AA",
                                     "csl": {"type": "article-journal", "title": "Human emotion",
                                             "author": [{"family": "Smith", "given": "J"}],
                                             "issued": {"date-parts": [[2007]]}},
                                     "doi": "", "first_author": "Smith", "year": "2007"}}
            cit_n, _ = zfi.insert_zotero_fields(src, out, (0, 1), [1], mapping, uri_prefix="http://zotero.org/users/U/items/")
            self.assertEqual(cit_n, 1)
            ok, det = zfi.verify(out, ["AA"], src_docx=src, item_mapping=mapping)
            self.assertEqual(det["missed_citations"], [])


class TestVerifyAllowsAbstract(unittest.TestCase):
    """真实 Zotero 域常内嵌 abstract，verify 不应据此判失败。"""

    def test_abstract_only_counts_informationally(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, "s.docx"); out = os.path.join(td, "o.docx")
            d = Document()
            d.add_paragraph("Claim (Smith, 2007).")
            d.add_paragraph("[1] Smith J. Human emotion. Cognition. 2007.")
            d.save(src)
            mapping = {"smith2007": {"itemKey": "AA", "uri": "u/AA",
                                     "csl": {"type": "article-journal", "title": "Human emotion",
                                             "author": [{"family": "Smith", "given": "J"}],
                                             "issued": {"date-parts": [[2007]]}},
                                     "doi": "", "first_author": "Smith", "year": "2007"}}
            zfi.insert_zotero_fields(src, out, (0, 1), [1], mapping, uri_prefix="http://zotero.org/users/U/items/")
            # 往 itemData 里手工塞一个 abstract，模拟真实 Zotero 域
            import zipfile
            from lxml import etree as ET
            z = zipfile.ZipFile(out)
            root = ET.fromstring(z.read('word/document.xml'))
            for it in root.iter(zfi.W + 'instrText'):
                if 'CSL_CITATION' in (it.text or ''):
                    obj = json.loads(it.text.split('CSL_CITATION', 1)[1].strip())
                    obj['citationItems'][0]['itemData']['abstract'] = 'Some abstract'
                    it.text = ' ADDIN ZOTERO_ITEM CSL_CITATION ' + json.dumps(obj, ensure_ascii=False) + ' '
                    break
            tmp = os.path.join(td, "o2.docx")
            with zipfile.ZipFile(out) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
                for f in zin.infolist():
                    zout.writestr(f, ET.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
                                  if f.filename == 'word/document.xml' else zin.read(f.filename))
            ok, det = zfi.verify(tmp, ["AA"])
            self.assertTrue(ok, det)
            self.assertEqual(det["incomplete"], 0)
            self.assertEqual(det["with_abstract"], 1)


class TestTableDoiSelfCell(unittest.TestCase):
    """refnum_column=None：引文域直接替换 DOI 格自身（整列即参考文献的表格）。"""

    def test_replace_doi_cell(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, "s.docx"); out = os.path.join(td, "o.docx")
            d = Document()
            d.add_paragraph("Intro")
            tbl = d.add_table(rows=2, cols=2)
            tbl.rows[0].cells[0].text = "Trial"
            tbl.rows[0].cells[1].text = "Reference"
            tbl.rows[1].cells[0].text = "Study A"
            tbl.rows[1].cells[1].text = "10.1234/ax;"
            d.save(src)
            mapping = {
                "ax": {"itemKey": "AAAAAAAA", "uri": "http://zotero.org/users/U/items/AAAAAAAA", "doi": "10.1234/ax",
                       "csl": {"type": "article-journal", "title": "A",
                               "author": [{"family": "A", "given": "X"}],
                               "issued": {"date-parts": [[2007]]}}}}
            n = zfi.insert_table_citations(src, out, mapping, 1, None, uri_prefix="http://zotero.org/users/U/items/")
            self.assertEqual(n, 1)
            ok, det = zfi.verify(out, ["AAAAAAAA"])
            self.assertTrue(ok, det)
            self.assertEqual(det["cites"], 1)


class TestVerifyZeroMissWithDoi(unittest.TestCase):
    """verify 的零遗漏校验必须复用与 insert 相同的 combined finder（作者-年份 + DOI 组
    + 文本括号 + 独立 DOI）。修复前只认作者-年份，DOI 风格正文被全部误报为漏掉。"""

    def _doi_mapping(self):
        return {
            "ax": {"itemKey": "AAAAAAAA", "uri": "http://zotero.org/users/U/items/AAAAAAAA", "doi": "10.1234/ax",
                   "csl": {"type": "article-journal", "title": "A", "author": [{"family": "A", "given": "X"}],
                           "issued": {"date-parts": [[2007]]}}},
            "fda": {"itemKey": "F1F1F1F1", "uri": "http://zotero.org/users/U/items/F1F1F1F1", "doi": "",
                    "csl": {"type": "report", "title": "FDA label", "author": [{"literal": "FDA"}],
                            "issued": {"date-parts": [[2020]]}}},
        }

    def _tm(self):
        return {"FDA label, minocycline": ["fda"]}

    def test_doi_style_doc_verify_zero_miss(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, "s.docx"); out = os.path.join(td, "o.docx")
            d = Document()
            d.add_paragraph("Claims (10.1234/ax) and (FDA label, minocycline) here.")
            d.add_paragraph("More (doi:10.1234/ax).")
            d.add_paragraph("[1] A. 2007.")
            d.save(src)
            m = self._doi_mapping(); tm = self._tm()
            cit_n, _ = zfi.insert_zotero_fields(
                src, out, (0, 2), None, m, uri_prefix="http://zotero.org/users/U/items/",
                citation_text_map=tm, include_tables=False)
            self.assertEqual(cit_n, 3)
            # 修复前：verify 只用作者-年份 -> 3 条 DOI/文本引文全部误报为 missed
            ok, det = zfi.verify(out, ["AAAAAAAA", "F1F1F1F1"],
                                 src_docx=src, item_mapping=m, citation_text_map=tm)
            self.assertTrue(ok, det)
            self.assertEqual(det["missed_citations"], [])
            self.assertEqual(det["display_mismatches"], [])

    def test_standalone_doi_cell_verify_zero_miss(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, "s.docx"); out = os.path.join(td, "o.docx")
            d = Document()
            d.add_paragraph("Intro")
            tbl = d.add_table(rows=2, cols=1)
            tbl.rows[0].cells[0].text = "Ref"
            tbl.rows[1].cells[0].text = "10.1234/ax;"
            d.save(src)
            m = self._doi_mapping()
            cit_n, _ = zfi.insert_zotero_fields(
                src, out, (0, 1), None, m, uri_prefix="http://zotero.org/users/U/items/",
                include_tables=True)
            self.assertEqual(cit_n, 1)
            ok, det = zfi.verify(out, ["AAAAAAAA"], src_docx=src, item_mapping=m)
            self.assertTrue(ok, det)
            self.assertEqual(det["missed_citations"], [])


class TestRtfUnescape(unittest.TestCase):
    """Zotero 在 formattedCitation 里用 RTF 十进制转义（\\u8211{} = en-dash）。"""

    def test_decimal_escapeseq_pred(self):
        BS = chr(92)
        self.assertEqual(zfi._rtf_unescape('[14,21' + BS + 'uc0' + BS + 'u8211' + '{}25]'),
                         '[14,21–25]')  # 十进制 8211 = 0x2013 = en-dash

    def test_norm_equal_after_unescape(self):
        BS = chr(92)
        f = '[1' + BS + 'uc0' + BS + 'u8211' + '{}3]'
        d = '[1–3]'
        self.assertEqual(zfi._norm(zfi._rtf_unescape(d)), zfi._norm(zfi._rtf_unescape(f)))


class TestOrgCreatorLiteral(unittest.TestCase):
    """以库为准重建时，organization 作者（firstName 空）应转成 CSL literal。"""

    def _fake(self):
        class Fake:
            def __call__(self, name, args=None):
                return {"key": args.get("itemKey"), "itemType": "report", "title": "FDA label",
                        "creators": [{"firstName": "", "lastName": "U.S. Food And Drug Administration",
                                      "creatorType": "author"}],
                        "date": "2020", "DOI": "", "url": ""}
        return Fake()

    def test_org_creator_literal(self):
        mapping = {"fdax": {"itemKey": "ABCD1234", "uri": "http://zotero.org/users/U/items/ABCD1234",
                            "csl": {"type": "report", "title": "FDA label"},
                            "doi": "", "first_author": "", "year": ""}}
        mp, mism = ri.apply_library_data(mapping, zotero=self._fake())
        self.assertEqual(mism, [])
        self.assertEqual(mp["fdax"]["csl"]["author"], [{"literal": "U.S. Food And Drug Administration"}])


class TestChineseSupport(unittest.TestCase):
    """中文文献：解析、引文匹配（含全角标点）、_norm 保留中文。"""

    def test_norm_keeps_chinese(self):
        self.assertEqual(ri._norm("张三"), "张三")
        self.assertEqual(ri.normalize_title("中文标题测试"), "中文标题测试")
        words = ri._title_words("儿童自闭症干预研究")
        self.assertIn("儿童", words)   # 中文二元组
        self.assertIn("自闭", words)

    def test_parse_chinese_reference(self):
        pr = ri.parse_numbered_reference("[1] 张三. 中文标题研究. 中华医学杂志. 2020.")
        self.assertEqual(pr["authors"][0]["family"], "张三")
        self.assertEqual(pr["title"], "中文标题研究")
        self.assertEqual(pr["journal"], "中华医学杂志")
        self.assertEqual(pr["year"], "2020")

    def test_match_chinese_citations(self):
        mapping = {"zhangsan2020": {"itemKey": "A1B2C3D4", "uri": "http://zotero.org/users/U/items/A1B2C3D4",
                                    "csl": {"type": "article-journal", "title": "中文标题研究",
                                            "author": [{"family": "张", "given": "三"}],
                                            "issued": {"date-parts": [[2020]]}},
                                    "doi": "10.x/1", "first_author": "张", "year": "2020"}}
        by_year, nih = zfi.build_matcher(mapping)
        for text in ["研究指出（张三，2020）。", "张三 (2020) 提出。", "参见（张三等，2020）。"]:
            spans = zfi.find_citations(text, by_year, nih)
            self.assertEqual(len(spans), 1, text)
            self.assertEqual(spans[0][3], ["zhangsan2020"])

    def test_chinese_end_to_end(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, "s.docx"); out = os.path.join(td, "o.docx")
            d = Document()
            d.add_paragraph("研究指出（张三，2020）。")
            d.add_paragraph("[1] 张三. 中文标题研究. 中华医学杂志. 2020.")
            d.save(src)
            mapping = {"zhangsan2020": {"itemKey": "A1B2C3D4", "uri": "http://zotero.org/users/U/items/A1B2C3D4",
                                        "csl": {"type": "article-journal", "title": "中文标题研究",
                                                "author": [{"family": "张", "given": "三"}],
                                                "issued": {"date-parts": [[2020]]}},
                                        "doi": "", "first_author": "张", "year": "2020"}}
            cit_n, _ = zfi.insert_zotero_fields(src, out, (0, 1), [1], mapping,
                                                uri_prefix="http://zotero.org/users/U/items/")
            self.assertEqual(cit_n, 1)
            ok, det = zfi.verify(out, ["A1B2C3D4"], src_docx=src, item_mapping=mapping)
            self.assertTrue(ok, det)
            self.assertEqual(det["cites"], 1)
            self.assertEqual(det["missed_citations"], [])


if __name__ == "__main__":
    unittest.main()